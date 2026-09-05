from decimal import Decimal as D
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from journal import Journal
from quotes import QuoteBook
from pnl import report
from export_pnl import export_report


class PnlTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.j=Journal(Path(self.tmp.name)/'journal.sqlite3')
        QuoteBook(self.j)

    def tearDown(self):
        self.j.close()
        self.tmp.cleanup()

    def trade(self,side,q,p,fee='0',day='2020-01-02',symbol='TEST'):
        return self.j.save(symbol=symbol,side=side,quantity=q,price=p,fee=fee,traded_on=day,market='US',currency='USD',note='=HYPERLINK("example")')

    def price(self,price='18',day='2020-01-03',symbol='TEST'):
        with self.j.db:
            self.j.db.execute('INSERT OR REPLACE INTO closes VALUES(?,?,?)',(symbol,day,price))

    def test_weighted_cost_fees_partial_sale_and_full_exit(self):
        self.trade('买入','10','10','2')
        self.trade('买入','10','20','2')
        self.trade('卖出','5','25','1')
        self.price()
        r=report(self.j,'2020-01-03')['rows'][0]
        self.assertEqual(r['cost'],D('228'))
        self.assertEqual(r['average'],D('15.2'))
        self.assertEqual(r['realized'],D('48'))
        self.assertEqual(r['unrealized'],D('42'))
        self.assertEqual(r['total'],D('90'))
        self.trade('卖出','15','18','3',day='2020-01-04')
        r=report(self.j,'2020-01-04')['rows'][0]
        self.assertEqual(r['cost'],0)
        self.assertEqual(r['total'],D('87'))
        self.assertEqual(r['unrealized'],0)
        self.trade('买入','1','30',day='2020-01-05')
        r=report(self.j,'2020-01-05')['rows'][0]
        self.assertEqual(r['average'],30)
        self.assertEqual(r['realized'],87)
        self.assertIsNone(r['total'])

    def test_date_alignment_and_missing_not_zero(self):
        self.trade('买入','10','10')
        self.price()
        self.trade('卖出','10','25',day='2020-01-04')
        r=report(self.j,'2020-01-03')
        self.assertEqual(r['rows'][0]['quantity'],10)
        self.assertEqual(r['totals']['realized'],0)
        self.assertEqual(r['excluded'],1)
        self.trade('买入','1','20',symbol='OTHER')
        r=report(self.j,'2020-01-03')
        self.assertIsNone(r['totals']['value'])
        self.assertIsNone(r['totals']['total'])
        self.assertEqual(r['totals']['realized'],0)

    def test_fee_validation_and_rounding_reconciliation(self):
        for fee in ('NaN','-1','Infinity','bad'):
            with self.assertRaises(ValueError):
                self.trade('买入','1','10',fee)
        self.trade('买入','3','1','1')
        self.trade('卖出','1','2')
        self.trade('卖出','2','2')
        self.assertEqual(report(self.j,'2020-01-03')['totals']['total'],2)

    def test_export_three_sheets_matches_snapshot_and_escapes_formulas(self):
        self.trade('买入','2','10','1')
        self.trade('卖出','1','20','1')
        data=report(self.j,'2020-01-03')
        path=Path(self.tmp.name)/'report.xlsx'
        export_report(data,path)
        wb=load_workbook(path,data_only=False)
        try:
            self.assertEqual(wb.sheetnames,['综合汇总','交易明细','计算说明'])
            self.assertEqual(wb['综合汇总']['H2'].value,8.5)
            self.assertIsNone(wb['综合汇总']['J2'].value)
            self.assertEqual(wb['交易明细']['O2'].data_type,'s')
            self.assertTrue(wb['交易明细']['O2'].value.startswith('='))
            self.assertEqual(wb['综合汇总'].freeze_panes,'A2')
        finally:
            wb.close()


if __name__=='__main__':
    unittest.main()
