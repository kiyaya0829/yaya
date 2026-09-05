"""Export an immutable report snapshot; workbook values mirror the app."""
from decimal import Decimal
from pathlib import Path
import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_report(data, destination):
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for values in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in values])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Never interpret user notes/symbols as Excel formulas.
                if isinstance(cell.value, (float,int)):
                    cell.number_format = '#,##0.00########;[Red](#,##0.00########);0.00'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True,color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='203047')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for i in range(1,len(headers)+1):
            ws.column_dimensions[get_column_letter(i)].width = 20
        ws.sheet_view.showGridLines = False
        return ws
    fields = ('symbol','quantity','average','cost','close','value','realized','unrealized','total','fees','status')
    rows = [[data['as_of'], *[r[k] for k in fields]] for r in data['rows']]
    total = data['totals']
    rows.append([data['as_of'],'合计',None,None,total['cost'],None,total['value'],total['realized'],total['unrealized'],total['total'],total['fees'],'缺价时合计留空' if total['total'] is None else '完整'])
    sheet('综合汇总',['估值日期','股票','当日持仓','平均成本 USD','剩余成本 USD','当日收盘 USD','收盘估值 USD','已实现盈亏 USD','浮动盈亏 USD','合计盈亏 USD','累计手续费 USD','状态'],rows)
    detail_fields = ('id','traded_on','market','symbol','side','currency','price','quantity','fee','released_cost','realized','remaining_quantity','remaining_cost','reason','note')
    sheet('交易明细',['记录 ID','日期','市场','股票','方向','币种','成交价','数量','手续费','结转成本 USD','本笔已实现 USD','本笔后持仓','本笔后成本 USD','是否计入','笔记'],[[Decimal(r[k]) if k in ('price','quantity','fee') else r[k] for k in detail_fields] for r in data['details']])
    notes = [
        ('估值日期',data['as_of']),('币种','US/USD 普通多头；其他市场和待确认记录不计入汇总'),
        ('日期对齐','只统计估值日及以前的交易；已实现、浮动、合计都截至同一日期。晚于该日期的交易仅列入明细。'),
        ('算法','移动加权平均。买入成本=数量×价格+手续费；部分卖出按卖出前平均成本结转；卖出收入扣手续费。清仓成本归零。'),
        ('费用','手续费已经计入盈亏；累计手续费仅供参考，不要再次扣除。旧版记录默认为 0，请补全实际费用。'),
        ('行情','Alpha Vantage 未复权日线；仅使用估值日的收盘价，不以其他日期替代。'),
        ('缺失值','持仓缺收盘价时，估值、浮动和合计为空；已实现盈亏仍可计算。清仓后不需要行情。'),
        ('范围限制','不处理拆股、分红、转仓、做空、税务。存在这些情况时需先校正记录，否则结果不适用。'),
        ('计算精度','应用使用 Decimal；导出为数值快照，Excel 数值精度有限（约15位有效数字），修改导出文件不会重新计算。'),
        ('排除记录数量',str(data['excluded'])),('资料','https://www.alphavantage.co/documentation/#daily')]
    ws = sheet('计算说明',['项目','说明'],notes)
    ws.column_dimensions['B'].width = 100
    for i in range(2,ws.max_row+1):
        ws.row_dimensions[i].height = 44
    destination = Path(destination)
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=destination.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary,destination)
    finally:
        wb.close()
        Path(temporary).unlink(missing_ok=True)
